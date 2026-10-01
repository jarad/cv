"""Build cv-style/reference.docx, the styles and page setup Word output uses.

Run from the repository root (needs `pip install python-docx` and pandoc):

    python3 cv-style/make-reference-docx.py

Word has no stylesheet language we can write by hand, so the look lives here
as code. Pandoc applies these styles by name; `cv.lua` and `R/cv.R` ask for
"CV Entry", "CV Group" and "CV New".
"""
import subprocess
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PATH = "cv-style/reference.docx"
NAME = "Jarad B. Niemi"
FONT = "Georgia"

with open(PATH, "wb") as f:
    f.write(subprocess.check_output(
        ["pandoc", "--print-default-data-file", "reference.docx"]))

d = Document(PATH)
styles = {s.name: s for s in d.styles}


def font(style, size=None, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def rule_below(style, sz):
    ppr = style.element.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    border.append(bottom)
    # schema order: pBdr comes early in pPr, before spacing and indentation
    ppr.insert(0, border)


for n in ("Normal", "Body Text", "First Paragraph", "Compact"):
    font(styles[n], size=10)
    styles[n].paragraph_format.space_before = Pt(0)
    styles[n].paragraph_format.space_after = Pt(3)
# table cells use Compact: keep a row's text together on one page
styles["Compact"].paragraph_format.keep_together = True

h1 = styles["Heading 1"]; font(h1, size=19, bold=True)
h1.paragraph_format.space_before = Pt(0); h1.paragraph_format.space_after = Pt(6)
h2 = styles["Heading 2"]; font(h2, size=11.5, bold=True)
h2.font.small_caps = True
h2.paragraph_format.space_before = Pt(14); h2.paragraph_format.space_after = Pt(5)
h3 = styles["Heading 3"]; font(h3, size=10.5, bold=True, italic=False)
h3.paragraph_format.space_before = Pt(6); h3.paragraph_format.space_after = Pt(4)
rule_below(h2, 4)
rule_below(h1, 8)

# a reference: hanging indent, a little air between entries, never split
entry = d.styles.add_style("CV Entry", WD_STYLE_TYPE.PARAGRAPH)
entry.base_style = styles["Normal"]
entry.paragraph_format.left_indent = Inches(0.2)
entry.paragraph_format.first_line_indent = Inches(-0.2)
entry.paragraph_format.space_after = Pt(7)
entry.paragraph_format.keep_together = True

# the label above a group of rows stays with them
group = d.styles.add_style("CV Group", WD_STYLE_TYPE.PARAGRAPH)
group.base_style = styles["Normal"]
group.font.bold = True
group.paragraph_format.space_before = Pt(6)
group.paragraph_format.keep_with_next = True

new = d.styles.add_style("CV New", WD_STYLE_TYPE.CHARACTER)
new.font.bold = True
new.font.size = Pt(7)
new.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)

# page setup, and a running header from page 2: name, then "Page n / total"
section = d.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)   # US letter
section.left_margin = section.right_margin = Inches(0.9)
section.top_margin = section.bottom_margin = Inches(0.8)
section.different_first_page_header_footer = True      # page 1 has no header


def field(run, code):
    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == code:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {code} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


para = section.header.paragraphs[0]
width = section.page_width - section.left_margin - section.right_margin
para.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
for text, code in ((NAME + "\tPage ", None), (None, "PAGE"), (" / ", None), (None, "NUMPAGES")):
    run = para.add_run(text or "")
    run.font.size = Pt(9)
    run.font.name = FONT
    if code:
        field(run, code)
ppr = para._p.get_or_add_pPr()
border = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "777777")):
    bottom.set(qn(k), v)
border.append(bottom)
ppr.insert(0, border)
section.first_page_header.is_linked_to_previous = False   # leave it empty

d.save(PATH)
