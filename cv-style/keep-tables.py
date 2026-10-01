"""Keep short tables on one page in the Word documents.

    python3 cv-style/keep-tables.py cv.docx cv-new.docx ...

Pandoc cannot ask Word to keep a table together, so this sets "keep with next"
on every paragraph of a table except its last row. A table of up to MAX_ROWS
rows stays whole and moves to the next page if it does not fit; a longer one
(the list of advisees, regular courses) is left to break, and Word repeats its
header row.
"""
import sys
from docx import Document

MAX_ROWS = 12

for path in sys.argv[1:]:
    doc = Document(path)
    for table in doc.tables:
        if len(table.rows) > MAX_ROWS:
            continue
        for row in table.rows[:-1]:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = True
    doc.save(path)
